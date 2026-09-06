"""PowerPoint extraction with one-based slide metadata."""

from pathlib import Path
from zipfile import BadZipFile

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from app.loaders.base import DocumentLoadError, DocumentRecord, clean_text, metadata_for, require_records, validate_document_path


def _shape_text(shape: object) -> list[str]:
    values: list[str] = []
    if getattr(shape, "has_text_frame", False):
        text = clean_text(getattr(shape, "text", ""))
        if text:
            values.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            row_text = " | ".join(clean_text(cell.text) for cell in row.cells if clean_text(cell.text))
            if row_text:
                values.append(row_text)
    return values


def load_pptx(path: str | Path) -> list[DocumentRecord]:
    resolved = validate_document_path(path, ".pptx")
    try:
        presentation = Presentation(resolved)
    except (PackageNotFoundError, BadZipFile, ValueError, KeyError) as exc:
        raise DocumentLoadError(f"Could not read PPTX {resolved.name}: the file may be corrupted.") from exc
    records: list[DocumentRecord] = []
    for index, slide in enumerate(presentation.slides):
        text = clean_text("\n".join(value for shape in slide.shapes for value in _shape_text(shape)))
        if text:
            records.append(DocumentRecord(text, metadata_for(resolved, "pptx", slide_number=index + 1)))
    return require_records(records, resolved)
