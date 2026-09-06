import tempfile
import unittest
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation

from app.loaders import DocumentLoadError
from app.rag.loader import load_document, supported_extensions


class DocumentLoaderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_txt_loader_decodes_utf16(self) -> None:
        path = self.root / "notes.txt"
        path.write_text("Deadlock prevention", encoding="utf-16")
        records = load_document(path)
        self.assertEqual(records[0].text, "Deadlock prevention")
        self.assertEqual(records[0].metadata, {"file_name": "notes.txt", "file_type": "txt"})

    def test_markdown_loader_preserves_headings(self) -> None:
        path = self.root / "notes.md"
        path.write_text("# DBMS\nNormalization notes\n## BCNF\nBCNF details", encoding="utf-8")
        records = load_document(path)
        self.assertEqual([record.metadata["section"] for record in records], ["DBMS", "BCNF"])
        self.assertIn("BCNF details", records[1].text)

    def test_pdf_loader_preserves_page_numbers(self) -> None:
        path = self.root / "course.pdf"
        document = pymupdf.open()
        document.new_page().insert_text((72, 72), "Page one")
        document.new_page().insert_text((72, 72), "Page two")
        document.save(path)
        document.close()
        records = load_document(path)
        self.assertEqual([record.metadata["page_number"] for record in records], [1, 2])
        self.assertEqual(records[1].metadata["file_type"], "pdf")

    def test_docx_loader_preserves_sections(self) -> None:
        path = self.root / "course.docx"
        document = Document()
        document.add_heading("Normalization", level=1)
        document.add_paragraph("First normal form")
        document.add_heading("Transactions", level=1)
        document.add_paragraph("Atomicity")
        document.save(path)
        records = load_document(path)
        self.assertEqual([record.metadata["section"] for record in records], ["Normalization", "Transactions"])
        self.assertIn("Atomicity", records[1].text)

    def test_pptx_loader_preserves_slide_numbers(self) -> None:
        path = self.root / "slides.pptx"
        presentation = Presentation()
        first = presentation.slides.add_slide(presentation.slide_layouts[1])
        first.shapes.title.text = "Networks"
        first.placeholders[1].text = "OSI model"
        second = presentation.slides.add_slide(presentation.slide_layouts[1])
        second.shapes.title.text = "Security"
        second.placeholders[1].text = "Encryption"
        presentation.save(path)
        records = load_document(path)
        self.assertEqual([record.metadata["slide_number"] for record in records], [1, 2])
        self.assertIn("Encryption", records[1].text)

    def test_unsupported_and_empty_documents_fail_safely(self) -> None:
        unsupported = self.root / "notes.csv"
        unsupported.write_text("a,b", encoding="utf-8")
        empty = self.root / "empty.txt"
        empty.write_text("", encoding="utf-8")
        with self.assertRaisesRegex(DocumentLoadError, "Unsupported document type"):
            load_document(unsupported)
        with self.assertRaisesRegex(DocumentLoadError, "No readable text"):
            load_document(empty)
        self.assertEqual(set(supported_extensions()), {".pdf", ".txt", ".md", ".docx", ".pptx"})

    def test_missing_and_corrupted_documents_fail_safely(self) -> None:
        with self.assertRaisesRegex(DocumentLoadError, "Document not found"):
            load_document(self.root / "missing.txt")
        for file_name in ("bad.pdf", "bad.docx", "bad.pptx"):
            path = self.root / file_name
            path.write_bytes(b"not a valid document")
            with self.subTest(file_name=file_name):
                with self.assertRaisesRegex(DocumentLoadError, "may be corrupted"):
                    load_document(path)


if __name__ == "__main__":
    unittest.main()
